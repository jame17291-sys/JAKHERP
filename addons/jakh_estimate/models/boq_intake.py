"""Reviewed, atomic XLSX intake. Files are evidence, never executable input."""
import hashlib
import io
import json
import zipfile
from uuid import uuid4

from openpyxl import load_workbook

from odoo import api, fields, models, _
from odoo.exceptions import AccessError, UserError, ValidationError

from .estimate import FINANCIAL, _window
from ..services.costing import calculate_estimate, json_value


DEFAULT_MAPPING = {'code': 'Code', 'description': 'Description', 'quantity': 'Quantity',
                   'unit_cost': 'Unit cost', 'uom': 'Unit', 'cost_code': 'Cost code',
                   'product': 'Product', 'parent_code': 'Parent code',
                   'allocation_percent': 'Parent %', 'is_group': 'Group'}


class BoqImportBatch(models.Model):
    _name = 'jakh.boq.import.batch'
    _description = 'Reviewed BOQ Workbook Intake'
    _inherit = ['mail.thread']
    _check_company_auto = True

    name = fields.Char(required=True, default=lambda self: self.env._('BOQ intake'))
    estimate_id = fields.Many2one('jakh.estimate', required=True, check_company=True, ondelete='restrict')
    company_id = fields.Many2one(related='estimate_id.company_id', store=True)
    currency_id = fields.Many2one(related='estimate_id.currency_id', groups=FINANCIAL)
    file_ids = fields.One2many('jakh.boq.import.file', 'batch_id', copy=False)
    row_ids = fields.One2many('jakh.boq.import.row', 'batch_id', readonly=True, copy=False)
    state = fields.Selection([('draft', 'Draft mapping'), ('preview', 'Previewed'), ('reviewed', 'Reviewed'), ('committed', 'Committed')], default='draft', required=True, readonly=True, tracking=True)
    preview_json = fields.Json(readonly=True, groups=FINANCIAL)
    preview_text = fields.Text(compute='_compute_preview_display', groups=FINANCIAL)
    error_count = fields.Integer(readonly=True)
    fingerprint = fields.Char(readonly=True, copy=False, groups=FINANCIAL)
    reviewed_by = fields.Many2one('res.users', readonly=True, copy=False)
    reviewed_at = fields.Datetime(readonly=True, copy=False)
    committed_at = fields.Datetime(readonly=True, copy=False)
    confirmed = fields.Boolean(string='I reconciled each file and approve these rows for the draft BOQ')

    _server_fields = {'state', 'row_ids', 'preview_text', 'preview_json', 'error_count', 'fingerprint', 'reviewed_by', 'reviewed_at', 'committed_at'}

    @api.depends('preview_json')
    @api.depends_context('lang')
    def _compute_preview_display(self):
        for batch in self:
            preview = batch.preview_json or {}
            lines = [self.env._('Company currency: %s', preview.get('company_currency', ''))] if preview else []
            for source in preview.get('files', []):
                lines.append(self.env._('Workbook %(file)s · %(rows)s rows · %(errors)s errors · quantity %(quantity)s · direct cost %(direct)s',
                    file=source['file'], rows=source['rows'], errors=source['errors'],
                    quantity=source.get('quantity', '—'), direct=source.get('direct_total', '—')))
            if preview.get('combined_direct_total') is not None:
                lines.append(self.env._('Combined BOQ: direct cost %(direct)s · cost budget %(budget)s · selling value %(selling)s',
                    direct=preview['combined_direct_total'], budget=preview['combined_budget_total'], selling=preview['combined_selling_total']))
            for error in preview.get('errors', []):
                lines.append(self.env._('%(file)s / %(sheet)s / row %(row)s: %(error)s', **error))
            batch.preview_text = '\n'.join(lines)

    def _locked(self):
        self.ensure_one(); self.check_access('write'); self.estimate_id._ensure_mutable()
        self.env.cr.execute('SELECT id FROM jakh_boq_import_batch WHERE id = %s FOR UPDATE', [self.id])
        self.invalidate_recordset()

    @api.model_create_multi
    def create(self, values_list):
        self.env['jakh.estimate']._check_editor()
        if any(set(values) & self._server_fields for values in values_list):
            raise AccessError(_('Intake results and workflow are assigned by reviewed actions.'))
        records = super().create(values_list)
        records.mapped('estimate_id')._ensure_mutable()
        return records

    def write(self, values):
        if set(values) & self._server_fields or 'estimate_id' in values:
            raise AccessError(_('Intake identity and workflow are server controlled.'))
        for row in self:
            row._locked()
            if row.state != 'draft' and not (row.state == 'preview' and set(values) == {'confirmed'}):
                raise UserError(_('Reset an uncommitted preview before editing its inputs.'))
        return super().write(values)

    def unlink(self):
        if any(row.state != 'draft' for row in self):
            raise UserError(_('Keep reviewed workbook evidence and its provenance.'))
        self.mapped('estimate_id')._ensure_mutable()
        return super().unlink()

    def _input_fingerprint(self):
        value = [{'id': row.id, 'sha256': hashlib.sha256(row.upload.content).hexdigest(),
                  'mapping': row._mapping(), 'sheet': row.sheet_name, 'header': row.header_row,
                  'product': row.default_product_id.id, 'cost_code': row.default_cost_code_id.id,
                  'expected': row.expected_direct_total, 'expect_zero': row.reconcile_zero_total} for row in self.file_ids.sorted('id')]
        value.append({'estimate': self.estimate_id._payload(), 'currency': self.currency_id.id})
        return hashlib.sha256(json.dumps(value, sort_keys=True, default=str).encode()).hexdigest()

    def action_reset(self):
        self._locked()
        if self.state == 'committed':
            raise UserError(_('A committed intake is retained permanently.'))
        self.row_ids._discard_preview()
        super().write({'state': 'draft', 'preview_json': False, 'error_count': 0, 'fingerprint': False,
                       'reviewed_by': False, 'reviewed_at': False, 'confirmed': False})
        return _window(self._name, self.id, _('Reviewed BOQ intake'))

    def action_preview(self):
        self._locked()
        if self.state != 'draft' or not self.file_ids:
            raise UserError(_('Add XLSX files to a draft intake before previewing.'))
        if len(self.file_ids) > 20 or sum(source.upload.size for source in self.file_ids) > 50 * 1024 * 1024:
            raise UserError(_('A reviewed intake supports up to 20 files and 50 MiB in total.'))
        parsed, errors, summaries = [], [], []
        for source in self.file_ids:
            rows, failures = source._parse_rows()
            parsed.extend(rows); errors.extend(failures)
            summaries.append({'id': source.id, 'file': source.filename, 'rows': len(rows), 'errors': len(failures)})
        if len(parsed) + len(self.estimate_id.line_ids) > 500:
            errors.append({'file': '', 'sheet': '', 'row': 0, 'error': _('An estimate supports up to 500 BOQ lines.')})
        identities = {line.code: line.line_uuid for line in self.estimate_id.line_ids}
        for row in parsed:
            if row['code'] in identities:
                errors.append({**row['provenance'], 'error': _('Duplicate BOQ code in the estimate or workbook batch: %s', row['code'])})
            else:
                identities[row['code']] = row['uuid']
        for row in parsed:
            if row.get('parent_code') and row['parent_code'] not in identities:
                errors.append({**row['provenance'], 'error': _('Unknown parent BOQ code: %s', row['parent_code'])})
            row['payload']['parent_uuid'] = identities.get(row.get('parent_code'), '')
        totals = None
        if not errors:
            payload = self.estimate_id._payload()
            payload['lines'] += [row['payload'] for row in parsed]
            try:
                totals = json_value(calculate_estimate(payload, translate=self.env._))
                by_id = {line['id']: line for line in totals['lines']}
                for source in self.file_ids:
                    own = [by_id[row['uuid']] for row in parsed if row['file_id'] == source.id and not row['payload']['is_group']]
                    direct = sum(float(row['cost_total']) for row in own)
                    summary = next(item for item in summaries if item['id'] == source.id)
                    summary.update(direct_total=direct, quantity=sum(float(row['total_quantity']) for row in own))
                    if (source.expected_direct_total or source.reconcile_zero_total) and self.currency_id.compare_amounts(direct, source.expected_direct_total):
                        errors.append({'file': source.filename, 'sheet': source.sheet_name or '', 'row': 0,
                                       'error': _('Workbook direct-cost total does not reconcile to the expected amount.')})
            except ValueError as exc:
                errors.append({'file': '', 'sheet': '', 'row': 0, 'error': str(exc)})
        values = [{'batch_id': self.id, 'file_id': row['file_id'], 'source_sheet': row['provenance']['sheet'],
                   'source_row': row['provenance']['row'], 'code': row['code'], 'line_uuid': row['uuid'],
                   'payload_json': json_value(row['payload']), 'values_json': row['values']} for row in parsed]
        self.env['jakh.boq.import.row']._create_preview(values)
        super().write({'state': 'preview', 'error_count': len(errors), 'fingerprint': self._input_fingerprint(),
                       'confirmed': False,
                       'preview_json': {'files': summaries, 'errors': errors,
                           'company_currency': self.currency_id.name, 'combined_direct_total': totals['cost_total'] if totals else None,
                           'combined_budget_total': totals['budget_total'] if totals else None,
                           'combined_selling_total': totals['sell_total'] if totals else None}})
        return _window(self._name, self.id, _('Reviewed BOQ intake'))

    def action_review(self):
        self._locked()
        if self.state != 'preview' or self.error_count or not self.row_ids or not self.confirmed:
            raise UserError(_('Review a clean, nonempty preview and confirm file reconciliation first.'))
        if self.fingerprint != self._input_fingerprint():
            raise UserError(_('The estimate or workbook inputs changed. Reset and preview again.'))
        super().write({'state': 'reviewed', 'reviewed_by': self.env.uid, 'reviewed_at': fields.Datetime.now()})
        self.env['jakh.audit.event']._append(self.company_id, 'boq.intake_reviewed', self, {'rows': len(self.row_ids), 'fingerprint': self.fingerprint})
        return True

    def action_commit(self):
        self.ensure_one(); self.check_access('write')
        if self.state == 'committed':
            return _window('jakh.estimate', self.estimate_id.id, _('BOQ estimate'))
        self._locked()
        if self.state == 'committed':
            return _window('jakh.estimate', self.estimate_id.id, _('BOQ estimate'))
        if self.state != 'reviewed' or self.error_count or self.fingerprint != self._input_fingerprint():
            raise UserError(_('Commit only the unchanged, clean, reviewed workbook batch.'))
        # One database savepoint covers all rows and the committed state.
        with self.env.cr.savepoint():
            values = [{**row.values_json, 'estimate_id': self.estimate_id.id, 'line_uuid': row.line_uuid,
                       'parent_uuid': row.payload_json.get('parent_uuid') or False, 'import_row_id': row.id} for row in self.row_ids]
            self.env['jakh.estimate.line']._create_from_intake(values, self)
            self.estimate_id._calculate()
            super().write({'state': 'committed', 'committed_at': fields.Datetime.now()})
            self.env['jakh.audit.event']._append(self.company_id, 'boq.intake_committed', self,
                                               {'rows': len(values), 'fingerprint': self.fingerprint})
        return _window('jakh.estimate', self.estimate_id.id, _('BOQ estimate'))


class BoqImportFile(models.Model):
    _name = 'jakh.boq.import.file'
    _description = 'BOQ Workbook Source'
    _check_company_auto = True
    batch_id = fields.Many2one('jakh.boq.import.batch', required=True, ondelete='cascade', check_company=True)
    company_id = fields.Many2one(related='batch_id.company_id', store=True)
    currency_id = fields.Many2one(related='batch_id.currency_id')
    upload = fields.Binary(required=True, attachment=False, groups=FINANCIAL)
    filename = fields.Char(required=True)
    sheet_name = fields.Char(string='Sheet (blank = first sheet)')
    header_row = fields.Integer(default=1, required=True)
    code_column = fields.Char(string='BOQ code column', default='Code', required=True)
    description_column = fields.Char(string='Description column', default='Description', required=True)
    quantity_column = fields.Char(string='Quantity column', default='Quantity', required=True)
    unit_cost_column = fields.Char(string='Direct unit cost column', default='Unit cost')
    uom_column = fields.Char(string='Unit column', default='Unit', required=True)
    cost_code_column = fields.Char(string='Cost-code column', default='Cost code')
    product_column = fields.Char(string='Product reference column', default='Product')
    parent_code_column = fields.Char(string='Parent BOQ code column', default='Parent code')
    allocation_percent_column = fields.Char(string='Parent share column', default='Parent %')
    is_group_column = fields.Char(string='Parent group flag column', default='Group')
    default_product_id = fields.Many2one('product.product', check_company=True)
    default_cost_code_id = fields.Many2one('jakh.cost.code', check_company=True)
    expected_direct_total = fields.Monetary(default=0, groups=FINANCIAL)
    reconcile_zero_total = fields.Boolean(string='Expected total is explicitly zero', default=False)

    def _mapping(self):
        self.ensure_one()
        return {key: self[key + '_column'] or '' for key in DEFAULT_MAPPING}

    def _editable(self):
        for row in self:
            row.batch_id._locked()
            if row.batch_id.state != 'draft':
                raise UserError(_('Reset the preview before changing files or mappings.'))

    @api.model_create_multi
    def create(self, values_list):
        batches = self.env['jakh.boq.import.batch'].browse([value['batch_id'] for value in values_list])
        for batch in batches:
            batch._locked()
            if batch.state != 'draft':
                raise UserError(_('Add source files only to a draft intake.'))
        return super().create(values_list)

    def write(self, values):
        if 'batch_id' in values:
            raise AccessError(_('Workbook source cannot move to another intake.'))
        self._editable(); return super().write(values)

    def unlink(self):
        self._editable(); return super().unlink()

    def _parse_rows(self):
        self.ensure_one()
        data = self.upload.content
        if len(data) > 20 * 1024 * 1024 or not self.filename.lower().endswith('.xlsx'):
            return [], [{'file': self.filename, 'sheet': '', 'row': 0, 'error': _('Use an XLSX file no larger than 20 MiB.')}]
        failures, rows = [], []
        book = None
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if len(archive.infolist()) > 5000 or sum(entry.file_size for entry in archive.infolist()) > 100 * 1024 * 1024:
                    raise ValueError(_('Expanded workbook exceeds the safe intake limit.'))
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            sheet = book[self.sheet_name] if self.sheet_name else book.worksheets[0]
            if not 1 <= self.header_row <= 100 or sheet.max_row > 10000 or sheet.max_column > 200:
                raise ValueError(_('Workbook dimensions or header row exceed the supported limits.'))
            headers = [str(cell.value).strip() if cell.value is not None else '' for cell in next(sheet.iter_rows(min_row=self.header_row, max_row=self.header_row))]
            mapping = self._mapping()
            if not isinstance(mapping, dict) or set(mapping) - set(DEFAULT_MAPPING):
                raise ValueError(_('Map only supported BOQ fields to workbook headers.'))
            positions = {}
            for target, header in mapping.items():
                if not isinstance(header, str):
                    raise ValueError(_('Column mapping values must be header text.'))
                if header and header in headers:
                    if headers.count(header) != 1:
                        raise ValueError(_('Mapped workbook headers must be unique.'))
                    positions[target] = headers.index(header)
            if not {'code', 'description', 'quantity', 'uom'}.issubset(positions):
                raise ValueError(_('Map code, description, quantity and unit columns before previewing.'))
            for number, cells in enumerate(sheet.iter_rows(min_row=self.header_row + 1), self.header_row + 1):
                if not any(cell.value is not None for cell in cells):
                    continue
                if len(rows) + len(failures) >= 500:
                    failures.append({'file': self.filename, 'sheet': sheet.title, 'row': number,
                                     'error': _('An estimate supports up to 500 BOQ lines.')})
                    break
                provenance = {'file': self.filename, 'sheet': sheet.title, 'row': number}
                try:
                    if any(cells[index].data_type == 'f' for index in positions.values()):
                        raise ValueError(_('Formula cells must be reviewed and supplied as values; formulas are never executed.'))
                    raw = {key: cells[index].value for key, index in positions.items()}
                    code = str(raw.get('code') or '').strip()
                    description = str(raw.get('description') or '').strip()
                    unit = str(raw.get('uom') or '').strip()
                    uoms = self.env['uom.uom'].with_context(lang='en_US').search([('name', '=', unit)], limit=2)
                    if not uoms:
                        uoms = self.env['uom.uom'].with_context(lang='ar_001').search([('name', '=', unit)], limit=2)
                    if len(uoms) != 1:
                        raise ValueError(_('The unit must match one unique existing ERP unit.'))
                    products = self.default_product_id
                    if raw.get('product'):
                        products = self.env['product.product'].search([('default_code', '=', str(raw['product']).strip()), ('company_id', 'in', [False, self.company_id.id])], limit=2)
                    codes = self.default_cost_code_id
                    if raw.get('cost_code'):
                        codes = self.env['jakh.cost.code'].search([('code', '=', str(raw['cost_code']).strip()), ('company_id', '=', self.company_id.id)], limit=2)
                    group = raw.get('is_group') in (True, 1, '1', 'yes', 'Yes', 'نعم')
                    if raw.get('is_group') not in (None, '', False, 0, '0', 'no', 'No', 'لا', True, 1, '1', 'yes', 'Yes', 'نعم'):
                        raise ValueError(_('Use yes/no for the parent group flag.'))
                    if not group and (len(products) != 1 or len(codes) != 1):
                        raise ValueError(_('Each leaf requires one existing product and company cost code.'))
                    if not group and uoms != products.uom_id:
                        raise ValueError(_('Imported unit must match the product primary unit; review conversions separately before handover.'))
                    uid = str(uuid4())
                    rate = 0 if raw.get('unit_cost') in (None, '') else raw['unit_cost']
                    share = 0 if raw.get('allocation_percent') in (None, '') else raw['allocation_percent']
                    payload = {'id': uid, 'code': code, 'description': description, 'quantity': raw.get('quantity'),
                               'unit_cost': rate, 'uom': uoms.name, 'cost_code': codes.code if len(codes) == 1 else '',
                               'product_id': products.id if len(products) == 1 else False, 'is_group': group,
                               'allocation_percent': share}
                    # Validate row inputs independently, then validate relationships as one batch.
                    checked = {**payload, 'allocation_percent': 0}
                    calculate_estimate({'lines': [checked]}, translate=self.env._)
                    values = {'code': code, 'name': description, 'quantity': raw.get('quantity'), 'unit_cost': rate,
                              'uom_id': uoms.id, 'product_id': products.id if len(products) == 1 else False,
                              'cost_code_id': codes.id if len(codes) == 1 else False, 'is_group': group,
                              'allocation_percent': share}
                    rows.append({'code': code, 'uuid': uid, 'file_id': self.id, 'parent_code': str(raw.get('parent_code') or '').strip(),
                                 'payload': payload, 'values': values, 'provenance': provenance})
                except (ValueError, ValidationError) as exc:
                    failures.append({**provenance, 'error': str(exc)})
        except (ValueError, KeyError, zipfile.BadZipFile, OSError) as exc:
            failures.append({'file': self.filename, 'sheet': self.sheet_name or '', 'row': 0, 'error': str(exc)})
        finally:
            if book:
                book.close()
        return rows, failures


class BoqImportRow(models.Model):
    _name = 'jakh.boq.import.row'
    _description = 'Immutable BOQ Import Provenance'
    _check_company_auto = True
    batch_id = fields.Many2one('jakh.boq.import.batch', required=True, check_company=True, ondelete='restrict')
    company_id = fields.Many2one(related='batch_id.company_id', store=True)
    file_id = fields.Many2one('jakh.boq.import.file', required=True, check_company=True, ondelete='restrict')
    source_sheet = fields.Char(required=True)
    source_row = fields.Integer(required=True)
    code = fields.Char(required=True)
    line_uuid = fields.Char(required=True)
    payload_json = fields.Json(required=True, groups=FINANCIAL)
    values_json = fields.Json(required=True, groups=FINANCIAL)
    _unique_source = models.Constraint('UNIQUE(file_id, source_sheet, source_row)', 'A source workbook row can appear only once in an intake.')

    @api.model_create_multi
    def create(self, values):
        raise AccessError(_('BOQ provenance is assigned by the reviewed preview.'))

    @api.model
    def _create_preview(self, values):
        return super().create(values)

    def write(self, values):
        raise AccessError(_('BOQ provenance cannot be rewritten.'))

    def unlink(self):
        raise AccessError(_('BOQ provenance is retained with its source workbook.'))

    def _discard_preview(self):
        if any(row.batch_id.state == 'committed' for row in self):
            raise UserError(_('Committed BOQ provenance cannot be deleted.'))
        return super().unlink()
